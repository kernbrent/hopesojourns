"""Build the user-facing MMT Word guide from its canonical Markdown source."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from build_tech_admin_guides import (
    COLORS,
    HEADING_FONT,
    LOGO_PATH,
    add_inline_text,
    add_running_furniture,
    configure_page,
    configure_styles,
    create_numbering,
    extract_metadata,
    mark_document_update_fields,
    render_markdown,
    set_run_font,
)
from build_mmt_user_guide_web import build as build_web_edition


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs/user/Hope-Sojourns-MMT-User-Guide.md"
OUTPUT = ROOT / "docs/user/Hope-Sojourns-MMT-User-Guide.docx"


def slug(text: str) -> str:
    return re.sub(r"[^a-z0-9 -]", "", text.lower()).strip().replace(" ", "-")


def add_bookmark(paragraph, name: str, number: int) -> None:
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(number))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(number))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def add_internal_link(paragraph, label: str, bookmark: str, size: float | None = None) -> None:
    link = OxmlElement("w:hyperlink")
    link.set(qn("w:anchor"), bookmark)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), COLORS["forest"])
    props.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(underline)
    if size is not None:
        font_size = OxmlElement("w:sz")
        font_size.set(qn("w:val"), str(round(size * 2)))
        props.append(font_size)
    run.append(props)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    link.append(run)
    paragraph._p.append(link)


def add_page_reference(paragraph, bookmark: str) -> None:
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), f"PAGEREF {bookmark} \\h")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = ""
    run.append(text)
    field.append(run)
    paragraph._p.append(field)


def add_cover(doc: Document, version: str, reviewed: str) -> None:
    for _ in range(2):
        doc.add_paragraph()
    logo = doc.add_paragraph()
    logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    logo.add_run().add_picture(str(LOGO_PATH), width=Inches(5.4))
    logo.paragraph_format.space_after = Pt(26)
    label = doc.add_paragraph()
    label.alignment = WD_ALIGN_PARAGRAPH.CENTER
    label.paragraph_format.space_after = Pt(12)
    set_run_font(label.add_run("MINISTRY MANAGEMENT"), size=10, color=COLORS["gold_ink"], bold=True)
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(15)
    set_run_font(title.add_run("Hope Sojourns\nUser Guide"), name=HEADING_FONT, size=29, color=COLORS["forest_deep"], bold=True)
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(sub.add_run("Step-by-step instructions for the private MMT Portal"), size=12, color=COLORS["forest"])
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_before = Pt(25)
    set_run_font(meta.add_run(f"Version {version}  •  Last reviewed {reviewed}"), size=10, color=COLORS["muted"])
    footer = doc.add_paragraph()
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.space_before = Pt(20)
    set_run_font(footer.add_run("Christian Steps Ministries  |  Private staff reference"), size=9.5, color=COLORS["muted"])
    footer.add_run().add_break(WD_BREAK.PAGE)


def add_contents(doc: Document, headings: list[str]) -> None:
    title = doc.add_paragraph("Table of contents", style="Heading 1")
    title.paragraph_format.space_before = Pt(0)
    for heading in headings:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.tab_stops.add_tab_stop(Inches(6.65), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        add_internal_link(p, heading, slug(heading))
        p.add_run("\t")
        add_page_reference(p, slug(heading))
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    add_inline_text(p, "In desktop Word, Ctrl+click a section name to jump there; on a phone, tap it. The alphabetical index is at the end of the guide.", base_size=9.5)
    p.add_run().add_break(WD_BREAK.PAGE)


def add_index(doc: Document, lines: list[str], heading_lookup: dict[str, str]) -> None:
    heading = doc.add_paragraph("Index", style="Heading 1")
    heading.paragraph_format.page_break_before = True
    add_bookmark(heading, "index", 999)
    note = doc.add_paragraph()
    note.paragraph_format.space_after = Pt(6)
    add_inline_text(note, "Numbers refer to guide sections. In desktop Word, Ctrl+click a topic to jump there; on a phone, tap it.", base_size=9)
    entries = []
    for line in lines:
        match = re.match(r"^([A-Z]):\s*(.+)$", line)
        if match:
            entries.append((match.group(1), re.findall(r"\[([^]]+)\]\(#([^)]+)\)", match.group(2))))
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    for cell in table.rows[0].cells:
        cell.width = Inches(3.2)
    midpoint = (len(entries) + 1) // 2
    for col, group in enumerate((entries[:midpoint], entries[midpoint:])):
        cell = table.cell(0, col)
        for entry_number, (letter, items) in enumerate(group):
            p = cell.paragraphs[0] if entry_number == 0 else cell.add_paragraph()
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.05
            set_run_font(p.add_run(letter + "  "), size=10, color=COLORS["forest_deep"], bold=True)
            for i, (term, anchor) in enumerate(items):
                if i:
                    set_run_font(p.add_run("  ·  "), size=10, color=COLORS["ink"])
                add_internal_link(p, term, anchor, size=10)
                section = heading_lookup.get(anchor, "")
                if section:
                    set_run_font(p.add_run(" " + section), size=10, color=COLORS["ink"])


def build() -> None:
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    version, reviewed = extract_metadata(lines)
    toc_start = next(i for i, line in enumerate(lines) if line == "## Table of contents")
    body_start = next(i for i, line in enumerate(lines) if line.startswith("## 1 Start and navigate"))
    index_start = next(i for i, line in enumerate(lines) if line == "## Index")
    revision_start = next(i for i, line in enumerate(lines) if line == "## Revision history")
    headings = [m.group(1) for line in lines[body_start:index_start] if (m := re.match(r"##\s+(.+)$", line))]
    heading_lookup = {slug(h): h.split(" ", 1)[0] for h in headings}

    doc = Document()
    configure_page(doc.sections[0])
    configure_styles(doc)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.18
    doc.styles["Normal"].paragraph_format.space_after = Pt(5)
    for name in ("List Bullet", "List Number"):
        doc.styles[name].paragraph_format.line_spacing = 1.18
        doc.styles[name].paragraph_format.space_after = Pt(3)
    add_running_furniture(doc.sections[0], "MMT User Guide")
    doc.sections[0].footer.paragraphs[0].runs[0].text = "Hope Sojourns MMT user guide  •  "
    mark_document_update_fields(doc)
    doc.core_properties.title = "Hope Sojourns Ministry Management Portal User Guide"
    doc.core_properties.subject = "Step-by-step staff instructions for the Hope Sojourns MMT Portal"
    doc.core_properties.author = "Christian Steps Ministries"
    doc.core_properties.keywords = "Hope Sojourns, MMT, user guide, ministry management"
    doc.core_properties.comments = "Generated from docs/user/Hope-Sojourns-MMT-User-Guide.md"
    numbering = {kind: create_numbering(doc, kind) for kind in ("bullet", "number", "checkbox")}

    add_cover(doc, version, reviewed)
    add_contents(doc, headings + ["Index"])
    introduction = doc.add_paragraph()
    introduction.paragraph_format.space_after = Pt(10)
    add_inline_text(introduction, "This private operating guide is for authorized Hope Sojourns and Christian Steps Ministries staff. Available actions depend on your permissions. Live and test records are separate; confirm the environment before entering real information.", base_size=10)
    render_markdown(doc, lines[body_start:index_start], numbering)
    for n, heading in enumerate(headings, 1):
        paragraph = next((p for p in doc.paragraphs if p.style.name == "Heading 1" and p.text == heading), None)
        if paragraph is None:
            raise ValueError(f"Missing rendered heading: {heading}")
        add_bookmark(paragraph, slug(heading), n)
    add_index(doc, lines[index_start + 1:revision_start], heading_lookup)
    revision = next((line[2:] for line in lines[revision_start + 1:] if line.startswith("- ")), "")
    if revision:
        note = doc.add_paragraph()
        note.paragraph_format.space_before = Pt(8)
        note.paragraph_format.space_after = Pt(0)
        add_inline_text(note, "Revision history: " + revision, base_size=8.5)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    build_web_edition()
    print(OUTPUT)


if __name__ == "__main__":
    build()
